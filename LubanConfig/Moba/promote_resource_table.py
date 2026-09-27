"""Promote an Excel-authored resource document into a typed Luban table."""

import argparse
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from migrate_json_to_excel import TARGET, RESOURCE_ROOT, cell_value, column_type


TABLE_KEYS = {"effects": "Id", "brains": "BrainId"}
FLOAT_FIELDS = {"effects": {"BaseDamage", "AttackRatio", "Radius"}}
DOCUMENTS = TARGET / "Datas/resource_documents.xlsx"
CATALOG = TARGET / "Datas/__tables__.xlsx"
COMPLEX = TARGET / "complex_columns.json"
KEYS = TARGET / "table_keys.json"


def promote(name):
    key = TABLE_KEYS[name]
    table_path = TARGET / "Datas" / f"{name}.xlsx"
    if table_path.exists():
        raise ValueError(f"{table_path} already exists; edit it directly")

    document_book = load_workbook(DOCUMENTS)
    document_sheet = document_book.active
    resource_path = f"moba/{name}.json"
    chunks = []
    removal_rows = []
    for index in range(4, document_sheet.max_row + 1):
        if document_sheet.cell(index, 3).value == resource_path:
            chunks.append((document_sheet.cell(index, 4).value,
                           document_sheet.cell(index, 5).value))
            removal_rows.append(index)
    if not chunks:
        raise ValueError(f"{resource_path} is not present in {DOCUMENTS}")
    chunks.sort()
    if [part for part, _ in chunks] != list(range(len(chunks))):
        raise ValueError(f"{resource_path} has missing or duplicate parts")
    rows = json.loads("".join(value for _, value in chunks))
    if not isinstance(rows, list) or not rows or any(not isinstance(row, dict) for row in rows):
        raise ValueError(f"{resource_path} must contain a nonempty array of objects")
    if any(not isinstance(row.get(key), int) or row[key] <= 0 for row in rows):
        raise ValueError(f"{resource_path} requires positive {key} values")
    ids = [row[key] for row in rows]
    if len(ids) != len(set(ids)):
        raise ValueError(f"{resource_path} contains duplicate {key} values")

    published = RESOURCE_ROOT / resource_path
    if json.loads(published.read_text(encoding="utf-8-sig")) != rows:
        raise ValueError(f"{published} differs from the Excel source; reconcile before promotion")

    fields = [key, *sorted({field for row in rows for field in row if field != key})]
    if name == "brains":
        fields.append("SkillSelectionPolicy")
    kinds = {field: column_type([row.get(field) for row in rows]) for field in fields}
    for field in FLOAT_FIELDS.get(name, set()):
        kinds[field] = "float"
    complex_fields = {
        field: "array" if any(isinstance(row.get(field), list) for row in rows) else "object"
        for field in fields
        if any(isinstance(row.get(field), (list, dict)) for row in rows)
        and kinds[field] == "string"
    }

    book = Workbook()
    sheet = book.active
    sheet.title = name
    sheet.append(["##var", *fields, "OmittedFields", "NullFields"])
    sheet.append(["##type", *(kinds[field] for field in fields), "string", "string"])
    sheet.append(["##", "Promoted from resource_documents.xlsx; edit this table as the source."])
    for row in rows:
        sheet.append([None, *(cell_value(row.get(field), kinds[field]) for field in fields),
                      ",".join(field for field in fields if field not in row),
                      ",".join(field for field in fields if field in row and row[field] is None)])
    for cells in sheet:
        for cell in cells:
            cell.font = Font(name="Arial", size=10, bold=cell.row <= 2)
    sheet.freeze_panes = "B4"
    sheet.column_dimensions["A"].width = 12
    for index in range(2, len(fields) + 4):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = 22

    catalog_book = load_workbook(CATALOG)
    catalog_sheet = catalog_book.active
    catalog_sheet.append([None, name.title().replace("_", ""),
                          "DR" + name.title().replace("_", ""), True,
                          f"{name}@{name}.xlsx", key, None, "c,s,e", None, None, name])
    metadata = json.loads(COMPLEX.read_text(encoding="utf-8"))
    metadata[name] = complex_fields
    keys = json.loads(KEYS.read_text(encoding="utf-8")) if KEYS.exists() else {}
    keys[name] = key

    for index in reversed(removal_rows):
        document_sheet.delete_rows(index)
    book.save(table_path)
    document_book.save(DOCUMENTS)
    catalog_book.save(CATALOG)
    COMPLEX.write_text(json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    KEYS.write_text(json.dumps(keys, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"promoted {name}: {len(rows)} rows, key={key}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--table", required=True, choices=TABLE_KEYS)
    promote(parser.parse_args().table)
