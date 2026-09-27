"""One-time migration of MOBA runtime JSON arrays into Luban workbooks.

Existing MiniTemplate workbooks are intentionally left untouched. Once a table
has been migrated, rerunning without --replace refuses to overwrite Excel.
"""

import argparse
import json
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "Unity/Packages/com.abilitykit.demo.moba.view.runtime/Resources/moba"
RESOURCE_ROOT = SOURCE.parent
TEMPLATE = ROOT / "LubanConfig/Moba/MiniTemplate"
TARGET = ROOT / "LubanConfig/Moba/Production"
TABLES = (
    "characters", "attribute_templates", "attr_types", "skills",
    "passive_skills", "skill_flows", "skill_level_tables", "models",
    "buffs", "continuous_processes", "projectile_launchers", "projectiles",
    "aoes", "emitters", "summons", "summon_attr_inherits",
    "component_templates", "skill_button_templates", "tag_templates",
    "continuous_tag_templates", "search_query_templates",
    "spawn_summon_action_templates", "presentation_templates",
    "battle_maps", "gameplays", "motion_groups",
)
RESOURCE_TABLE = "resource_documents"
CHUNK_SIZE = 12000


def column_type(values):
    present = [value for value in values if value is not None]
    if not present:
        return "string"
    if all(isinstance(value, bool) for value in present):
        return "bool"
    if all(isinstance(value, int) and not isinstance(value, bool) for value in present):
        return "int"
    if all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in present):
        return "float"
    if all(isinstance(value, str) for value in present):
        return "string"
    if all(isinstance(value, list) and all(isinstance(item, int) for item in value) for value in present):
        return "(list#sep=,),int"
    if all(isinstance(value, list) and all(isinstance(item, str) and "," not in item for item in value) for value in present):
        return "(list#sep=,),string"
    return "string"


def cell_value(value, kind):
    if value is None:
        return ""
    if kind.startswith("(list#sep=,)"):
        return ",".join(str(item) for item in value)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def write_table(name, replace):
    path = TARGET / "Datas" / f"{name}.xlsx"
    if path.exists() and not replace:
        return
    rows = json.loads((SOURCE / f"{name}.json").read_text(encoding="utf-8-sig"))
    if not isinstance(rows, list):
        raise ValueError(f"{name}: expected a JSON array")
    fields = sorted({key for row in rows for key in row}) or ["Id"]
    if "Id" not in fields:
        raise ValueError(f"{name}: no Id field")
    fields.remove("Id")
    fields.insert(0, "Id")
    kinds = {key: column_type([row.get(key) for row in rows]) for key in fields}
    source_fields = fields[:]
    fields.extend(("OmittedFields", "NullFields"))
    kinds["OmittedFields"] = "string"
    kinds["NullFields"] = "string"
    book = Workbook()
    sheet = book.active
    sheet.title = name
    sheet.append(["##var", *fields])
    sheet.append(["##type", *(kinds[key] for key in fields)])
    sheet.append(["##", "Generated once from the runtime JSON baseline; edit Excel after migration."])
    for row in rows:
        values = [cell_value(row.get(key), kinds[key]) for key in source_fields]
        values.append(",".join(key for key in source_fields if key not in row))
        values.append(",".join(key for key in source_fields if key in row and row[key] is None))
        sheet.append([None, *values])
    for cells in sheet:
        for cell in cells:
            cell.font = Font(name="Arial", size=10, bold=cell.row <= 2)
    sheet.freeze_panes = "B4"
    sheet.column_dimensions["A"].width = 12
    for index in range(2, len(fields) + 2):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = 22
    book.save(path)
    print(f"migrated {name}: {len(rows)} rows, {len(fields)} fields")


def write_catalog(names, replace):
    destination = TARGET / "Datas/__tables__.xlsx"
    if destination.exists() and not replace:
        return
    source = TEMPLATE / "Datas/__tables__.xlsx"
    book = load_workbook(source)
    sheet = book.active
    if sheet.max_row > 3:
        sheet.delete_rows(4, sheet.max_row - 3)
    for name in names:
        sheet.append([None, name.title().replace("_", ""),
                      "DR" + name.title().replace("_", ""), True,
                      f"{name}@{name}.xlsx", "Id", None, "c,s,e", None, None, name])
    book.save(destination)


def write_complex_columns(names, replace):
    destination = TARGET / "complex_columns.json"
    if destination.exists() and not replace:
        return
    metadata = {}
    for name in names:
        rows = json.loads((SOURCE / f"{name}.json").read_text(encoding="utf-8-sig"))
        metadata[name] = {
            key: "array" if any(isinstance(row.get(key), list) for row in rows) else "object"
            for key in {field for row in rows for field in row}
            if any(isinstance(row.get(key), (list, dict)) for row in rows)
            and column_type([row.get(key) for row in rows]) == "string"
        }
    destination.write_text(
        json.dumps(metadata, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def write_resource_documents(replace):
    path = TARGET / "Datas" / f"{RESOURCE_TABLE}.xlsx"
    if path.exists() and not replace:
        return
    book = Workbook()
    sheet = book.active
    sheet.title = RESOURCE_TABLE
    sheet.append(["##var", "Id", "Path", "Part", "Content"])
    sheet.append(["##type", "int", "string", "int", "string"])
    sheet.append(["##", "Other MOBA resources; each Path is relative to Resources."])
    table_names = set(TABLES)
    files = sorted(RESOURCE_ROOT.joinpath("ability").rglob("*.json"))
    files += sorted(path for path in SOURCE.glob("*.json") if path.stem not in table_names)
    row_id = 1
    for file in files:
        relative = file.relative_to(RESOURCE_ROOT).as_posix()
        content = file.read_text(encoding="utf-8-sig")
        json.loads(content)
        for part, start in enumerate(range(0, len(content), CHUNK_SIZE)):
            sheet.append([None, row_id, relative, part, content[start:start + CHUNK_SIZE]])
            row_id += 1
    sheet.freeze_panes = "B4"
    sheet.column_dimensions["B"].width = 10
    sheet.column_dimensions["C"].width = 75
    sheet.column_dimensions["D"].width = 10
    sheet.column_dimensions["E"].width = 100
    book.save(path)
    print(f"migrated {len(files)} resource documents into {row_id - 1} chunks")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--table", choices=TABLES)
    parser.add_argument("--replace", action="store_true")
    parser.add_argument("--resources-only", action="store_true")
    args = parser.parse_args()
    if args.replace and (TARGET / "table_keys.json").exists():
        raise ValueError("Production has promoted tables; --replace would overwrite Excel sources")
    (TARGET / "Datas").mkdir(parents=True, exist_ok=True)
    (TARGET / "Defines").mkdir(exist_ok=True)
    for filename in ("__beans__.xlsx", "__enums__.xlsx"):
        destination = TARGET / "Datas" / filename
        if not destination.exists():
            shutil.copy2(TEMPLATE / "Datas" / filename, destination)
    builtin = TARGET / "Defines/builtin.xml"
    if not builtin.exists():
        shutil.copy2(TEMPLATE / "Defines/builtin.xml", builtin)
    config_path = TARGET / "luban.conf"
    if not config_path.exists():
        configuration = json.loads((TEMPLATE / "luban.conf").read_text(encoding="utf-8"))
        for target in configuration["targets"]:
            target["topModule"] = "moba_luban"
        config_path.write_text(
            json.dumps(configuration, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    names = TABLES[:TABLES.index(args.table) + 1] if args.table else TABLES
    if not args.resources_only:
        for name in names:
            write_table(name, args.replace)
        write_complex_columns(names, args.replace)
    write_resource_documents(args.replace)
    write_catalog((*names, RESOURCE_TABLE), args.replace)


if __name__ == "__main__":
    main()
