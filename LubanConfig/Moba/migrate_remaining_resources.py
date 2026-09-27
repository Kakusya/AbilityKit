"""One-time migration from resource_documents.xlsx to typed MOBA tables and Ability nodes."""

import json
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from migrate_json_to_excel import ROOT, TARGET, cell_value, column_type


DATAS = TARGET / "Datas"
DOCUMENTS = DATAS / "resource_documents.xlsx"
CATALOG = DATAS / "__tables__.xlsx"
COMPLEX = TARGET / "complex_columns.json"
KEYS = TARGET / "table_keys.json"
ROOTS = TARGET / "table_roots.json"
BACKUP = ROOT / "local/moba-config-migration/resource_documents.xlsx"
ABILITY_TABLE = "ability_nodes"
AGGREGATE = "ability/ability_trigger_plans.json"
MAX_VALUE = 12000


def read_documents():
    book = load_workbook(DOCUMENTS, read_only=True, data_only=True)
    chunks = {}
    try:
        for _, row_id, path, part, content in list(book.active.values)[3:]:
            if path is None:
                continue
            chunks.setdefault(path, []).append((part, content))
    finally:
        book.close()
    documents = {}
    for path, parts in chunks.items():
        parts.sort()
        if [number for number, _ in parts] != list(range(len(parts))):
            raise ValueError(f"{path}: invalid part sequence")
        documents[path] = json.loads("".join(content for _, content in parts))
    return documents


def write_workbook(name, rows, key, note):
    path = DATAS / f"{name}.xlsx"
    if path.exists():
        raise ValueError(f"{path} already exists")
    fields = [key, *sorted({field for row in rows for field in row if field != key})]
    kinds = {field: column_type([row.get(field) for row in rows]) for field in fields}
    complex_fields = {
        field: "array" if any(isinstance(row.get(field), list) for row in rows) else "object"
        for field in fields
        if any(isinstance(row.get(field), (list, dict)) for row in rows)
        and kinds[field] == "string"
    }
    book = Workbook()
    sheet = book.active
    sheet.title = name
    sheet.append(["##var", *fields, "OmittedFields", "NullFields"])
    sheet.append(["##type", *(kinds[field] for field in fields), "string", "string"])
    sheet.append(["##", note])
    for row in rows:
        sheet.append([None, *(cell_value(row.get(field), kinds[field]) for field in fields),
                      ",".join(field for field in fields if field not in row),
                      ",".join(field for field in fields if field in row and row[field] is None)])
    for cells in sheet:
        for cell in cells:
            cell.font = Font(name="Arial", size=10, bold=cell.row <= 2)
    sheet.freeze_panes = "B4"
    sheet.column_dimensions["A"].width = 12
    for index in range(2, len(fields) + 4):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = 24
    book.save(path)
    return complex_fields


def escape(token):
    return str(token).replace("~", "~0").replace("/", "~1")


def ability_rows(documents):
    rows = []

    def emit(path, pointer, value, split=False):
        encoded = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
        if split or len(encoded) > MAX_VALUE:
            if isinstance(value, dict):
                kind, children = "object", value.items()
            elif isinstance(value, list):
                kind, children = "array", enumerate(value)
            else:
                raise ValueError(f"{path}{pointer}: scalar exceeds Excel cell limit")
            rows.append({"Id": len(rows) + 1, "Path": path, "Pointer": pointer,
                         "Kind": kind, "ValueJson": ""})
            for token, child in children:
                emit(path, pointer + "/" + escape(token), child)
        else:
            rows.append({"Id": len(rows) + 1, "Path": path, "Pointer": pointer,
                         "Kind": "json", "ValueJson": "json:" + encoded})

    for path, value in sorted(documents.items()):
        if path.startswith("ability/") and path != AGGREGATE:
            emit(path, "", value, split=True)
    return rows


def main():
    if not DOCUMENTS.exists():
        raise ValueError("resource_documents.xlsx is already migrated")
    if ROOTS.exists() or (DATAS / f"{ABILITY_TABLE}.xlsx").exists():
        raise ValueError("remaining resources have already been migrated")
    documents = read_documents()
    if AGGREGATE not in documents:
        raise ValueError("the trigger aggregate baseline is missing")

    metadata = json.loads(COMPLEX.read_text(encoding="utf-8"))
    keys = json.loads(KEYS.read_text(encoding="utf-8"))
    roots = {}
    catalog_book = load_workbook(CATALOG)
    catalog_sheet = catalog_book.active
    for index in range(catalog_sheet.max_row, 3, -1):
        if catalog_sheet.cell(index, 11).value == "resource_documents":
            catalog_sheet.delete_rows(index)

    for path, source in sorted(documents.items()):
        if not path.startswith("moba/"):
            continue
        name = Path(path).stem
        rows = source if isinstance(source, list) else [source]
        if not rows or any(not isinstance(row, dict) for row in rows):
            raise ValueError(f"{path}: expected an object or a nonempty array of objects")
        key = next((candidate for candidate in ("Id", "id", "Code")
                    if all(isinstance(row.get(candidate), int) for row in rows)), "RowId")
        if key == "RowId":
            rows = [{"RowId": index, **row} for index, row in enumerate(rows, 1)]
        ids = [row[key] for row in rows]
        if len(ids) != len(set(ids)):
            raise ValueError(f"{path}: duplicate {key}")
        metadata[name] = write_workbook(name, rows, key, "MOBA runtime resource; edit Excel as the source.")
        keys[name] = key
        roots[name] = {"kind": "array" if isinstance(source, list) else "object",
                       "syntheticKey": key == "RowId"}
        type_name = "Resource" + name.title().replace("_", "")
        catalog_sheet.append([None, type_name,
                              "DR" + type_name, True,
                              f"{name}@{name}.xlsx", key, None, "c,s,e", None, None, name])
        print(f"migrated {path}: {len(rows)} rows")

    nodes = ability_rows(documents)
    write_workbook(ABILITY_TABLE, nodes, "Id", "Ability JSON tree nodes; Pointer uses RFC 6901.")
    metadata[ABILITY_TABLE] = {}
    catalog_sheet.append([None, "AbilityNodes", "DRAbilityNodes", True,
                          "ability_nodes@ability_nodes.xlsx", "Id", None, "c,s,e", None, None, ABILITY_TABLE])
    print(f"migrated ability resources: {len(nodes)} nodes")

    catalog_book.save(CATALOG)
    COMPLEX.write_text(json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    KEYS.write_text(json.dumps(keys, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    ROOTS.write_text(json.dumps(roots, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(DOCUMENTS), str(BACKUP))
    print(f"archived legacy document workbook at {BACKUP}")


if __name__ == "__main__":
    main()
